"""Build the offline OsteoPatch report and figures.

Requires Python packages already used by the document runtime: matplotlib and
python-docx. This script does not modify application dependencies or model data.
"""
from __future__ import annotations

import csv
import html
import json
import re
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIG = HERE / "figures"
SOURCE = HERE / "OsteoPatch_Full_Research_Report.md"
DOCX = HERE / "OsteoPatch_Full_Research_Report.docx"
HTML = HERE / "index.html"

INK = "#19343A"
TEAL = "#267C78"
RUST = "#C56B4C"
AMBER = "#D6A13B"
BLUE = "#5A7D9A"
PALE = "#F3F0E8"
LINE = "#D8D8D0"
MUTED = "#637277"
WHITE = "#FFFFFF"

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 10,
    "axes.titleweight": "bold",
    "axes.labelcolor": INK,
    "text.color": INK,
    "axes.edgecolor": LINE,
    "xtick.color": MUTED,
    "ytick.color": MUTED,
    "savefig.facecolor": WHITE,
    "figure.facecolor": WHITE,
})


def save(fig, name: str) -> None:
    fig.savefig(FIG / name, dpi=220, bbox_inches="tight", facecolor=WHITE)
    plt.close(fig)


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def make_figures() -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    g4 = json.loads((ROOT / "aidlc-docs/inception/model/g4/overall-oof-metrics.json").read_text())
    support = load_csv(ROOT / "aidlc-docs/inception/model/g4/group-class-support.csv")

    # Cohort funnel, with QC flags kept separate from exclusion counts.
    fig, ax = plt.subplots(figsize=(11.2, 4.5))
    stages = [
        ("Source inventory", 1144, INK),
        ("After mixed-label policy", 1091, BLUE),
        ("After content review", 1032, AMBER),
        ("Frozen eligible cohort", 1028, TEAL),
    ]
    x = np.arange(len(stages))
    vals = [v for _, v, _ in stages]
    colors = [c for _, _, c in stages]
    bars = ax.bar(x, vals, color=colors, width=0.64)
    ax.set_ylim(0, 1320)
    ax.set_ylabel("Patch tiles")
    ax.set_xticks(x, [s[0] for s in stages])
    ax.set_title("From source inventory to frozen training cohort", loc="left", fontsize=16, pad=16)
    ax.grid(axis="y", color=LINE, linewidth=0.7)
    ax.set_axisbelow(True)
    for bar, value in zip(bars, vals):
        ax.text(bar.get_x() + bar.get_width()/2, value + 25, f"{value:,}", ha="center", weight="bold", fontsize=12)
    ax.text(0.01, -0.26, "Held out: 53 mixed + 59 review + 4 near-duplicate candidates. All 1,144 retrieved/decoded; 0 corrupt; 0 exact or pixel duplicates. QC signal counts can overlap.",
            transform=ax.transAxes, fontsize=8.5, color=MUTED)
    ax.spines[["top", "right", "left"]].set_visible(False)
    save(fig, "data_curation_flow.png")

    # Group-by-class support.
    groups = [r["group"] for r in support if r["group"] != "TOTAL"]
    labels = ["NON_TUMOR", "VIABLE_TUMOR", "NECROSIS"]
    colors = [BLUE, RUST, TEAL]
    data = np.array([[int(r[label]) for label in labels] for r in support if r["group"] != "TOTAL"])
    fig, ax = plt.subplots(figsize=(9.4, 5.2))
    left = np.zeros(len(groups))
    for i, label in enumerate(labels):
        ax.barh(groups, data[:, i], left=left, color=colors[i], label=label, height=0.56)
        for j, n in enumerate(data[:, i]):
            if n >= 16:
                ax.text(left[j] + n/2, j, str(n), ha="center", va="center", color=WHITE, fontsize=9, weight="bold")
        left += data[:, i]
    ax.set_title("Eligible class support is concentrated unevenly by group", loc="left", fontsize=15, pad=14)
    ax.set_xlabel("Eligible patch tiles")
    ax.grid(axis="x", color=LINE, linewidth=0.7)
    ax.set_axisbelow(True)
    ax.legend(ncol=1, frameon=False, loc="upper right")
    ax.text(0.01, -0.18, "Case-3 has 3 viable tiles; Case-48 has 2 necrosis tiles; P9 has no viable or necrosis tiles.",
            transform=ax.transAxes, fontsize=9, color=MUTED)
    ax.spines[["top", "right", "left"]].set_visible(False)
    save(fig, "group_class_support.png")

    # Per-class metrics.
    names = ["NON_TUMOR", "VIABLE_TUMOR", "NECROSIS"]
    metrics = g4["per_class"]
    vals = np.array([[metrics[n]["precision"], metrics[n]["recall"], metrics[n]["f1"]] for n in names])
    fig, ax = plt.subplots(figsize=(9.6, 5.2))
    x = np.arange(3)
    width = 0.22
    for i, (lab, color) in enumerate(zip(["Precision", "Recall", "F1"], [BLUE, RUST, TEAL])):
        rects = ax.bar(x + (i-1)*width, vals[:, i], width, label=lab, color=color)
        for j, (rect, value) in enumerate(zip(rects, vals[:, i])):
            ax.text(rect.get_x()+rect.get_width()/2, value + 0.022, f"{value:.2f}", ha="center", fontsize=8, color=INK)
    ax.set_ylim(0, 1.08)
    ax.set_xticks(x, names)
    ax.set_ylabel("Score")
    ax.set_title("G4 pooled OOF performance is class-asymmetric", loc="left", fontsize=15, pad=14)
    ax.grid(axis="y", color=LINE, linewidth=0.7)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, ncol=3, loc="upper center")
    ax.annotate("Recall 0.110", xy=(1, vals[1, 1]), xytext=(1.18, 0.28),
                arrowprops={"arrowstyle": "->", "color": RUST, "lw": 1.2},
                color=RUST, fontsize=10, weight="bold")
    ax.spines[["top", "right", "left"]].set_visible(False)
    save(fig, "g4_class_metrics.png")

    # Confusion matrix with row proportions in smaller type.
    cm = np.array(g4["confusion_matrix_rows_true_cols_pred"])
    row_pct = cm / cm.sum(axis=1, keepdims=True)
    fig, ax = plt.subplots(figsize=(7.2, 6.3))
    im = ax.imshow(cm, cmap="YlGnBu", vmin=0)
    ax.set_xticks(range(3), names, rotation=16, ha="right")
    ax.set_yticks(range(3), names)
    ax.set_xlabel("Predicted class")
    ax.set_ylabel("Reference class")
    ax.set_title("Pooled out-of-fold confusion matrix", loc="left", fontsize=15, pad=16)
    for i in range(3):
        for j in range(3):
            color = WHITE if cm[i, j] > cm.max()*0.45 else INK
            ax.text(j, i-0.08, f"{cm[i,j]}", ha="center", va="center", color=color, fontsize=14, weight="bold")
            ax.text(j, i+0.22, f"{row_pct[i,j]*100:.1f}%", ha="center", va="center", color=color, fontsize=8.5)
    cbar = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04)
    cbar.set_label("Tile count")
    ax.text(0.01, -0.18, "Viable tumor: 32/290 correct; 153 assigned NECROSIS; 105 assigned NON_TUMOR.",
            transform=ax.transAxes, fontsize=9, color=MUTED)
    save(fig, "g4_confusion_matrix.png")

    # Majority comparison. Metrics are separate panels to preserve units.
    fig, axes = plt.subplots(1, 2, figsize=(9.3, 4.4))
    comparisons = [
        ("Macro-F1", [0.213404, 0.562311]),
        ("Accuracy", [0.470817, 0.670233]),
    ]
    for ax, (title, values) in zip(axes, comparisons):
        bars = ax.bar(["Majority\nreference", "G4"], values, color=[BLUE, TEAL], width=0.55)
        ax.set_ylim(0, 0.78)
        ax.set_title(title, loc="left", fontsize=13, weight="bold")
        ax.grid(axis="y", color=LINE, linewidth=0.7)
        ax.set_axisbelow(True)
        for b, v in zip(bars, values):
            ax.text(b.get_x()+b.get_width()/2, v+0.025, f"{v:.3f}", ha="center", weight="bold")
        ax.spines[["top", "right", "left"]].set_visible(False)
    fig.suptitle("Descriptive comparison on the same pooled G4 OOF set", x=0.08, ha="left", fontsize=15, weight="bold")
    fig.text(0.08, -0.02, "Not a significance test or evidence of external generalization.", color=MUTED, fontsize=9)
    save(fig, "g4_baseline_comparison.png")

    make_journey()
    make_integrated_architecture()


def box(ax, x, y, w, h, title, body, color=TEAL, title_size=9.3, body_size=7.8, fill=WHITE):
    patch = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.35,rounding_size=1.1",
                           linewidth=1.2, edgecolor=color, facecolor=fill)
    ax.add_patch(patch)
    ax.text(x+w/2, y+h-2.5, title, ha="center", va="top", fontsize=title_size, weight="bold", color=color)
    ax.text(x+w/2, y+h-6.3, body, ha="center", va="top", fontsize=body_size, linespacing=1.35, color=INK)


def arrow(ax, start, end, color=TEAL, style="-", lw=1.6, rad=0):
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=11,
                                 linewidth=lw, color=color, linestyle=style,
                                 connectionstyle=f"arc3,rad={rad}"))


def make_journey() -> None:
    fig, ax = plt.subplots(figsize=(12, 4.2))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 30)
    ax.axis("off")
    steps = [
        (2, "1  Sign in + project", "Local prototype identity\nmembership and scope", BLUE),
        (22, "2  Review queue", "Filters + search\nsmall margin first", TEAL),
        (42, "3  Inspect patch", "H&E pixels + 3 scores\nmargin, entropy, caveats", RUST),
        (62, "4  Human decision", "Accept / correct / defer\nrevision + idempotency", AMBER),
        (82, "5  Evidence out", "Append-only history\nCSV / JSON export", BLUE),
    ]
    for x, title, body, color in steps:
        box(ax, x, 8, 15, 13, title, body, color, 9.0, 7.6, PALE)
    for i in range(len(steps)-1):
        arrow(ax, (steps[i][0]+15.5, 14.5), (steps[i+1][0]-0.5, 14.5), INK)
    ax.text(2, 26, "A review journey that keeps model output and human action separate", fontsize=15, weight="bold", color=INK)
    ax.text(42, 3, "Corrections append a new event; the original prediction remains immutable. Export access is capability-gated.",
            fontsize=9, color=MUTED)
    save(fig, "reviewer_journey.png")


def make_integrated_architecture() -> None:
    fig, ax = plt.subplots(figsize=(15.5, 10.5))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")
    ax.text(2, 97, "OSTEOPATCH  |  TRAINING EVIDENCE + REVIEW PRODUCT", fontsize=18, weight="bold", color=INK)
    ax.text(2, 92.7, "Solid arrows show the principal data path. The two lanes are connected by explicit model identity and evidence boundaries.",
            fontsize=9.5, color=MUTED)

    # Upper research lane.
    ax.add_patch(FancyBboxPatch((1, 59), 98, 29, boxstyle="round,pad=0.5,rounding_size=1.3",
                                edgecolor=LINE, facecolor="#F7F6F1", linewidth=1.0))
    ax.text(2.5, 85.5, "A   DATA + AI DEVELOPMENT LANE", fontsize=10, weight="bold", color=BLUE)
    xs = [2, 18.2, 34.4, 50.6, 66.8, 83.0]
    titles = [
        "Source collection",
        "Ingestion + identity",
        "Label + QC gate",
        "Frozen cohort",
        "G4 training",
        "OOF evidence",
    ]
    bodies = [
        "TCIA H&E tiles\n1,144 patches\nmetadata join",
        "retrieve + TIFF check\nhash + decode\n1024² RGB 8-bit",
        "3 tissue classes\n53 mixed excluded\n59 review + 4 dup held",
        "1,028 eligible\n4 case/slide groups\npatient link unverified",
        "MobileNetV3-Small\nencoder frozen\nlinear 1024 → 3",
        "4 LOGO folds\nmacro-F1 0.562\nviable recall 0.110",
    ]
    colors = [BLUE, BLUE, AMBER, TEAL, TEAL, RUST]
    for x, title, body, color in zip(xs, titles, bodies, colors):
        box(ax, x, 66, 13.5, 15.5, title, body, color, 8.8, 7.2, WHITE)
    for i in range(5):
        arrow(ax, (xs[i]+13.8, 73.8), (xs[i+1]-0.4, 73.8), BLUE, lw=1.45)

    # Separate controlled research status banner.
    ax.add_patch(FancyBboxPatch((2, 48.5), 96, 6.7, boxstyle="round,pad=0.38,rounding_size=1.0",
                                edgecolor=RUST, facecolor="#FBF1EC", linewidth=1.2))
    ax.text(4, 52.2, "MODEL IDENTITY + QUALIFICATION BOUNDARY", fontsize=9.2, weight="bold", color=RUST)
    ax.text(4, 49.9,
            "Original G4 binary is absent. Live inference uses a separately hashed recovered head. The 8 Oct all-data artifact failed its gate and is not selectable.",
            fontsize=8.4, color=INK)
    arrow(ax, (89.5, 66), (89.5, 55.5), RUST, style="--", lw=1.4)

    # Lower product lane.
    ax.add_patch(FancyBboxPatch((1, 8), 98, 36.5, boxstyle="round,pad=0.5,rounding_size=1.3",
                                edgecolor=LINE, facecolor="#F7F6F1", linewidth=1.0))
    ax.text(2.5, 41.5, "B   LOCAL APPLICATION + REVIEW LANE", fontsize=10, weight="bold", color=TEAL)
    low_x = [2, 18.2, 34.4, 50.6, 66.8, 83.0]
    low_titles = [
        "User + project",
        "React reviewer UI",
        "FastAPI boundary",
        "Core review API",
        "SQLite records",
        "Optional model ops",
    ]
    low_bodies = [
        "6 local personas\ncapability + scope\ngates",
        "gallery + search\npatch / analysis\nimport + review",
        "JWT dev issuer\nPydantic contracts\nsame-origin routes",
        "shared G6 handlers\nimage + prediction\nreview + export",
        "QC metadata\nprediction immutable\nreview append-only",
        "recovered head\nlive_run / live_tile\ncontrastive Grad-CAM",
    ]
    low_colors = [BLUE, TEAL, TEAL, TEAL, AMBER, RUST]
    for x, title, body, color in zip(low_x, low_titles, low_bodies, low_colors):
        box(ax, x, 19.5, 13.5, 16, title, body, color, 8.8, 7.2, WHITE)
    for i in range(4):
        arrow(ax, (low_x[i]+13.8, 27.5), (low_x[i+1]-0.4, 27.5), TEAL, lw=1.45)
    # Optional upload/inference stays outside the normal review-write path.
    arrow(ax, (89.5, 19), (73.4, 17), AMBER, style="--", lw=1.3, rad=-0.08)
    ax.text(82, 12.5, "authorized upload scores stay in live tables", ha="center", fontsize=7.3, color=MUTED)
    ax.text(49, 11, "human action → append-only revision; prediction remains immutable", ha="center", fontsize=7.4, color=BLUE)

    ax.text(2, 4.5,
            "Local prototype. Auth, storage, governance, and demo checks do not imply production identity, external validation, or clinical readiness.",
            fontsize=8.6, color=MUTED)
    save(fig, "integrated_training_product_architecture.png")


def inline_html(text: str) -> str:
    text = html.escape(text, quote=False)
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"\*(.+?)\*", r"<em>\1</em>", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', text)
    text = re.sub(r"(?<![\"=])(https?://[^\s<]+)", r'<a href="\1">\1</a>', text)
    return text


def is_table_line(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [x.strip() for x in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows


def render_html(md: str) -> str:
    lines = md.splitlines()
    body: list[str] = []
    i = 0
    list_open = False
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            if list_open:
                body.append("</ul>")
                list_open = False
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            if list_open:
                body.append("</ul>")
                list_open = False
            level = len(m.group(1))
            title = m.group(2)
            slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")
            body.append(f'<h{level} id="{slug}">{inline_html(title)}</h{level}>')
            i += 1
            continue
        if line.startswith(">"):
            body.append(f'<blockquote>{inline_html(line.lstrip("> "))}</blockquote>')
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if match:
                alt, src = match.groups()
                body.append(f'<figure><img src="{html.escape(src)}" alt="{html.escape(alt)}"><figcaption>{inline_html(alt)}</figcaption></figure>')
            i += 1
            continue
        if is_table_line(line):
            block = []
            while i < len(lines) and is_table_line(lines[i]):
                block.append(lines[i])
                i += 1
            rows = table_rows(block)
            if rows:
                body.append("<table><thead><tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in rows[0]) + "</tr></thead><tbody>")
                for row in rows[1:]:
                    body.append("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in row) + "</tr>")
                body.append("</tbody></table>")
            continue
        if re.match(r"^\d+\.\s+", line):
            if list_open:
                body.append("</ul>")
                list_open = False
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i]))
                i += 1
            body.append("<ol>" + "".join(f"<li>{inline_html(x)}</li>" for x in items) + "</ol>")
            continue
        if line.startswith("- "):
            if not list_open:
                body.append("<ul>")
                list_open = True
            body.append(f"<li>{inline_html(line[2:])}</li>")
            i += 1
            continue
        if line.strip() == "---":
            body.append("<hr>")
            i += 1
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s+|>\s?|!\[|- |\d+\.\s+)", lines[i]) and not is_table_line(lines[i]) and lines[i].strip() != "---":
            para.append(lines[i].strip())
            i += 1
        body.append(f"<p>{inline_html(' '.join(para))}</p>")
    return """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>OsteoPatch — Full Research Report</title>
<style>
@page{size:letter;margin:0.7in 0.72in 0.75in}
:root{--ink:#19343a;--teal:#267c78;--rust:#c56b4c;--line:#d8d8d0;--paper:#f3f0e8;--muted:#637277}
*{box-sizing:border-box}html{background:#e9e7e0}body{margin:0 auto;max-width:980px;background:white;color:var(--ink);padding:52px 74px 80px;font:16px/1.58 Georgia,'Times New Roman',serif}
h1,h2,h3,h4{font-family:Arial,sans-serif;line-height:1.2;color:var(--ink)}h1{font-size:34px;letter-spacing:-.7px;border-bottom:2px solid var(--teal);padding-bottom:15px;margin:5px 0 28px}
h2{font-size:23px;margin:46px 0 16px;padding-top:12px;border-top:1px solid var(--line);break-before:page}
h3{font-size:17px;margin:28px 0 10px;color:var(--teal)}h4{font-size:15px}
p{margin:0 0 14px}strong{color:#18333a}a{color:#1c706c;overflow-wrap:anywhere}
blockquote{margin:18px 0 22px;padding:14px 18px;background:var(--paper);border-left:4px solid var(--rust);font-family:Arial,sans-serif;font-size:13px}
ul,ol{padding-left:25px;margin:8px 0 18px}li{padding-left:4px;margin:5px 0}
table{border-collapse:collapse;width:100%;margin:17px 0 23px;font:12px/1.38 Arial,sans-serif;break-inside:avoid}
th{background:var(--ink);color:white;text-align:left;font-weight:600}th,td{border:1px solid var(--line);padding:7px 8px;vertical-align:top}tr:nth-child(even) td{background:#f7f6f1}
figure{margin:24px 0 28px;break-inside:avoid;text-align:center}figure img{width:100%;height:auto;max-height:8.7in;object-fit:contain}figcaption{font:12px/1.4 Arial,sans-serif;color:var(--muted);margin:8px auto 0;max-width:95%;text-align:left}
hr{border:0;border-top:1px solid var(--line);margin:30px 0}.cover{min-height:8.8in;display:flex;flex-direction:column;justify-content:center;break-after:page}
@media screen{body{box-shadow:0 2px 30px #0001;margin:28px auto}}
@media print{html,body{background:white;box-shadow:none}body{max-width:none;padding:0}h2{break-before:page}figure img{max-height:8.6in}a{color:inherit;text-decoration:none}}
h1:first-child{margin-top:100px;font-size:34px;break-after:avoid}
h1:first-child+p{font:600 14px/1.5 Arial,sans-serif;color:var(--teal)}
h1:first-child+p+p,h1:first-child+p+p+p{font:13px Arial,sans-serif;color:var(--muted);margin-top:-8px}
</style></head><body><main>""" + "".join(body) + """</main></body></html>"""


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_inline_runs(paragraph, text: str) -> None:
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run(token)
        if token.startswith("**") and token.endswith("**"):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run.text = token[1:-1]
            run.italic = True
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(35, 51, 55)


def shade_cell(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color.lstrip("#"))
    tc_pr.append(shd)


def set_cell_text(cell, value: str, header: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_inline_runs(p, value)
    for run in p.runs:
        run.font.size = Pt(8.2)
        if header:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_docx_table(doc: Document, raw_rows: list[str]) -> None:
    rows = table_rows(raw_rows)
    if not rows:
        return
    ncols = max(map(len, rows))
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncols):
            value = row[j] if j < len(row) else ""
            set_cell_text(table.cell(i, j), value, header=(i == 0))
            if i == 0:
                shade_cell(table.cell(i, j), INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(md: str) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color in [
        ("Heading 1", 21, INK), ("Heading 2", 15, TEAL), ("Heading 3", 12, BLUE)
    ]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        style.paragraph_format.space_before = Pt(15)
        style.paragraph_format.space_after = Pt(6)

    header = sec.header.paragraphs[0]
    header.text = "OSTEOPATCH  /  RESEARCH REPORT"
    header.style = styles["Caption"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = sec.footer.paragraphs[0]
    footer.text = "Evidence snapshot  •  8 October 2026                                      Page "
    footer.style = styles["Caption"]
    add_page_field(footer)

    lines = md.splitlines()
    title = lines[0].removeprefix("# ").strip()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(105)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))
    for line in lines[2:5]:
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            add_inline_runs(p, line.replace("**", ""))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(155)
    p.add_run("A project-grounded account of the medical rationale, model-development evidence, software system, and research limitations.")
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.color.rgb = RGBColor.from_string(TEAL.lstrip("#"))
    doc.add_page_break()

    # Static contents page for portability and predictable rendering.
    doc.add_heading("Contents", level=1)
    toc = [
        "Abstract", "Introduction and research problem", "Medical background and motivation",
        "Related literature", "Product definition and intended use", "Data provenance and curation",
        "AI development and evaluation methods", "Model results", "Software engineering and system behavior",
        "Verification and benchmarking", "Discussion", "Limitations, ethics, and governance",
        "Future work", "Conclusion", "Appendices A–F", "References",
    ]
    for item in toc:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)
    doc.add_page_break()

    i = 0
    # Skip document title and the cover metadata lines; body begins at Abstract.
    while i < len(lines) and lines[i].strip() != "## Abstract":
        i += 1
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        heading = re.match(r"^(#{2,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 2 and (re.match(r"\d+\.", text) or text.startswith("Appendix") or text == "References"):
                doc.add_page_break()
            doc.add_heading(text, level=1 if level == 2 else min(level-1, 3))
            i += 1
            continue
        if line.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if m:
                alt, src = m.groups()
                image_path = (HERE / src).resolve()
                if image_path.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(image_path), width=Inches(6.35))
                    cap = doc.add_paragraph(alt, style="Caption")
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if is_table_line(line):
            block = []
            while i < len(lines) and is_table_line(lines[i]):
                block.append(lines[i])
                i += 1
            add_docx_table(doc, block)
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            add_inline_runs(p, line.lstrip("> "))
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, line[2:])
            i += 1
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{2,4}\s+|>\s?|!\[|- |\d+\.\s+)", lines[i]) and not is_table_line(lines[i]) and lines[i].strip() != "---":
            para.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        add_inline_runs(p, " ".join(para))
    doc.save(DOCX)


def main() -> None:
    make_figures()
    md = SOURCE.read_text(encoding="utf-8")
    HTML.write_text(render_html(md), encoding="utf-8")
    build_docx(md)
    print(f"Created {len(list(FIG.glob('*.png')))} figures")
    print(f"Markdown: {SOURCE}")
    print(f"HTML: {HTML}")
    print(f"DOCX: {DOCX}")


if __name__ == "__main__":
    main()
